"""Acceptance test 8 and friends, run against the real workbook."""
import datetime as dt
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula

from conftest import event
from hbs_scoreboard import etl, excel_writer, metrics
from hbs_scoreboard.models import Stage

FIXTURE = Path(__file__).parent / "fixtures" / "HBS_Weekly_Scoreboard.xlsx"

# Every payments / revenue cell in the live workbook, with its value.
PAYMENT_CELLS = {
    "Tune-Up Factory": ["E13", "F13", "E24", "F24", "E38", "F38", "E49", "F49",
                        "E60", "F60", "F71", "E85", "F86", "E98", "F98"],
    "Direct Install": ["E12", "F12", "E10", "F10"],
    "Company Results": ["E38", "F38", "E11", "F11"],
    "Recurring Revenue": ["E7", "F7", "E15", "F15"],
}


@pytest.fixture
def workbook(tmp_path):
    dst = tmp_path / "wb.xlsx"
    shutil.copyfile(FIXTURE, dst)
    return dst


@pytest.fixture
def scoreboard(cfg):
    events = [
        event(1, "100 Main St", Stage.PA_SUBMITTED, "BGE", "2026-09-02", 10000),
        event(2, "100 Main St", Stage.PA_SUBMITTED, "BGE", "2026-09-02", 10000),  # dupe
        event(3, "200 Oak Ave", Stage.PA_SUBMITTED, "BGE", "2026-09-03", 5000),
        event(4, "300 Pine Rd", Stage.PA_SUBMITTED, "HVAC BTU", "2026-09-03", 3840),
        event(5, "400 Cedar Ln", Stage.AUDIT, "Pepco", "2026-09-04", None),
        event(6, "500 Elm St", Stage.CO_RECEIVED, "BGE", "2026-09-05", 99999),
    ]
    return metrics.compute(etl.build(events, cfg), cfg, as_of=dt.date(2026, 9, 4))


# -- acceptance test 8 --------------------------------------------------------
def test_writer_leaves_every_payments_and_revenue_cell_untouched(workbook, scoreboard):
    before = {}
    wb = load_workbook(workbook)
    for sheet, cells in PAYMENT_CELLS.items():
        for c in cells:
            before[(sheet, c)] = wb[sheet][c].value
    wb.close()

    report = excel_writer.write(scoreboard, workbook, week_index=1)
    assert report.saved

    wb = load_workbook(workbook)
    for (sheet, c), value in before.items():
        assert wb[sheet][c].value == value, f"{sheet}!{c} was modified"
    wb.close()


def test_writer_refuses_to_touch_manual_rows(workbook, scoreboard):
    report = excel_writer.write(scoreboard, workbook, week_index=1)
    manual = [s for s in report.skipped if s.get("reason") == "manually owned — never written"]
    assert any("Payments Received" in s["kpi"] for s in manual)
    assert any("Revenue" in s["kpi"] for s in manual)
    assert any("Utilization" in s["kpi"] for s in manual)
    written_cells = {(s["sheet"], s["cell"]) for s in report.written}
    for sheet, cells in PAYMENT_CELLS.items():
        for c in cells:
            assert (sheet, c) not in written_cells


# -- formulas and formatting --------------------------------------------------
def test_writer_preserves_every_formula(workbook, scoreboard):
    def formulas(path):
        wb = load_workbook(path)
        out = {}
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    v = cell.value
                    v = v.text if isinstance(v, ArrayFormula) else v
                    if isinstance(v, str) and v.startswith("="):
                        out[f"{ws.title}!{cell.coordinate}"] = v
        wb.close()
        return out

    before = formulas(workbook)
    excel_writer.write(scoreboard, workbook, week_index=1)
    after = formulas(workbook)
    assert before == after, "formulas changed"
    assert len(before) > 300


def test_writer_preserves_formatting(workbook, scoreboard):
    def styles(path):
        wb = load_workbook(path)
        out = {}
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    out[f"{ws.title}!{cell.coordinate}"] = (
                        cell.number_format, cell.font.b, cell.font.color.rgb
                        if cell.font.color else None, cell.fill.fgColor.rgb)
        wb.close()
        return out

    before = styles(workbook)
    excel_writer.write(scoreboard, workbook, week_index=1)
    assert styles(workbook) == before


def test_writer_fills_the_expected_week_column(workbook, scoreboard):
    report = excel_writer.write(scoreboard, workbook, week_index=1)
    assert report.target_column == "E"
    wb = load_workbook(workbook)
    ws = wb["Tune-Up Factory"]
    # BGE Tune-Up PA submitted: 2 unique projects (one same-day duplicate removed)
    assert ws["E18"].value == 2
    assert ws["E19"].value == 15000.0
    # Prescriptive PA submitted
    assert ws["E7"].value == 1
    assert ws["E8"].value == 3840.0
    wb.close()


def test_no_data_leaves_cell_untouched_rather_than_writing_zero(workbook, scoreboard):
    wb = load_workbook(workbook)
    before = wb["Tune-Up Factory"]["E53"].value      # Delmarva audits, no data
    wb.close()
    excel_writer.write(scoreboard, workbook, week_index=1)
    wb = load_workbook(workbook)
    assert wb["Tune-Up Factory"]["E53"].value == before
    wb.close()
    # and it is reported as skipped, with the reason
    report = excel_writer.write(scoreboard, workbook, week_index=1, dry_run=True)
    assert any(s["cell"] == "E53" and "no data" in s["reason"] for s in report.skipped)


def test_dry_run_changes_nothing(workbook, scoreboard):
    before = workbook.read_bytes()
    report = excel_writer.write(scoreboard, workbook, week_index=1, dry_run=True)
    assert not report.saved
    assert workbook.read_bytes() == before
    assert report.written, "dry run should still report what it would write"


def test_reports_the_workbooks_existing_broken_formulas(workbook, scoreboard):
    report = excel_writer.write(scoreboard, workbook, week_index=1, dry_run=True)
    refs = {b["cell"] for b in report.broken_formulas if b["sheet"] == "Company Results"}
    # Row 44 of Company Results carries #REF! from a deleted row.
    assert {"E44", "G44"} <= refs


def test_detects_workbook_week_convention():
    info = excel_writer.detect_week_convention(FIXTURE)
    assert info["weeks_in_month"] == 4
    assert "9/7" in str(info["current_week"])
    assert info["headers"][0].endswith("9/7")
